"""Test PPTX normalization using yakumono_trans logic.

Processes the English and Japanese test PPTX files using the same
normalization logic as en.py and ja.py, saves the results.
"""

import os
import re
import sys
import unicodedata
from pptx import Presentation

# ── helpers (copied from en.py / ja.py) ──


def zen_to_han(char):
    cp = ord(char)
    if 0xFF01 <= cp <= 0xFF5E:
        return chr(cp - 0xFEE0)
    if cp == 0x3000:
        return ' '
    return char

def normalize_url_email(text):
    url_pattern = r'[\uFF48\uFF28hH][\uFF54\uFF34tT][\uFF54\uFF34tT][\uFF50\uFF30pP][\uFF53\uFF33sS]?[\uFF1A：:][\uFF0F／/][\uFF0F／/][^\s\u3000]+'
    email_pattern = r'[\w\uFF21-\uFF3A\uFF41-\uFF5A\uFF10-\uFF19][\w.\-\uFF21-\uFF3A\uFF41-\uFF5A\uFF10-\uFF19\uFF0E\uFF0D]*[\uFF20@][\w\uFF21-\uFF3A\uFF41-\uFF5A\uFF10-\uFF19][\w.\-\uFF21-\uFF3A\uFF41-\uFF5A\uFF10-\uFF19\uFF0E\uFF0D]*'
    combined = f'({url_pattern})|({email_pattern})'
    def replace_match(m):
        return ''.join(zen_to_han(c) for c in m.group(0))
    return re.sub(combined, replace_match, text)



# ── EN replacement dict ──
en_replacement = {
    '\u3000': '\u0020',
    '\u30FB': '\u2022',
    '\uFF01': '\u0021', '\uFF02': '\u0022', '\uFF03': '\u0023',
    '\uFF04': '\u0024', '\uFF05': '\u0025', '\uFF06': '\u0026',
    '\uFF07': '\u0027', '\uFF08': '\u0028', '\uFF09': '\u0029',
    '\uFF0A': '\u002A', '\uFF0B': '\u002B', '\uFF0C': '\u002C',
    '\uFF0D': '\u002D', '\uFF0E': '\u002E', '\uFF0F': '\u002F',
    '\uFF10': '\u0030', '\uFF11': '\u0031', '\uFF12': '\u0032',
    '\uFF13': '\u0033', '\uFF14': '\u0034', '\uFF15': '\u0035',
    '\uFF16': '\u0036', '\uFF17': '\u0037', '\uFF18': '\u0038',
    '\uFF19': '\u0039', '\uFF1A': '\u003A', '\uFF1B': '\u003B',
    '\uFF1C': '\u003C', '\uFF1D': '\u003D', '\uFF1E': '\u003E',
    '\uFF1F': '\u003F', '\uFF20': '\u0040',
    '\uFF21': '\u0041', '\uFF22': '\u0042', '\uFF23': '\u0043',
    '\uFF24': '\u0044', '\uFF25': '\u0045', '\uFF26': '\u0046',
    '\uFF27': '\u0047', '\uFF28': '\u0048', '\uFF29': '\u0049',
    '\uFF2A': '\u004A', '\uFF2B': '\u004B', '\uFF2C': '\u004C',
    '\uFF2D': '\u004D', '\uFF2E': '\u004E', '\uFF2F': '\u004F',
    '\uFF30': '\u0050', '\uFF31': '\u0051', '\uFF32': '\u0052',
    '\uFF33': '\u0053', '\uFF34': '\u0054', '\uFF35': '\u0055',
    '\uFF36': '\u0056', '\uFF37': '\u0057', '\uFF38': '\u0058',
    '\uFF39': '\u0059', '\uFF3A': '\u005A',
    '\uFF3B': '\u005B', '\uFF3C': '\u005C', '\uFF3D': '\u005D',
    '\uFF3E': '\u005E', '\uFF3F': '\u005F', '\uFF40': '\u0060',
    '\uFF41': '\u0061', '\uFF42': '\u0062', '\uFF43': '\u0063',
    '\uFF44': '\u0064', '\uFF45': '\u0065', '\uFF46': '\u0066',
    '\uFF47': '\u0067', '\uFF48': '\u0068', '\uFF49': '\u0069',
    '\uFF4A': '\u006A', '\uFF4B': '\u006B', '\uFF4C': '\u006C',
    '\uFF4D': '\u006D', '\uFF4E': '\u006E', '\uFF4F': '\u006F',
    '\uFF50': '\u0070', '\uFF51': '\u0071', '\uFF52': '\u0072',
    '\uFF53': '\u0073', '\uFF54': '\u0074', '\uFF55': '\u0075',
    '\uFF56': '\u0076', '\uFF57': '\u0077', '\uFF58': '\u0078',
    '\uFF59': '\u0079', '\uFF5A': '\u007A',
    '\uFF5B': '\u007B', '\uFF5C': '\u007C', '\uFF5D': '\u007D',
    '\uFF5E': '\u007E',
    '\u33A1': 'm2', '\u33A5': 'm3',
    '\\\\ ': '¥', '("': '("', '\\\\': '¥',
    "'": "'", '\" ': '"\\\\s', ' \"': '\\\\s"',
    '")': '")', '\".': '".', '.\"': '."',
    '・\\t': '• ', ' )': ')', ' . ': '. ',
    ' , ': ', ', '\" ': '」 ',
    '( ': '(', "' ": "' ", "'": "'",
    "• ": "•   ", '） ': '）', ' （': '（',
    '.\"': '."', "' ": "' ",
    ' 、': '、', '。 ': '。',
    ' 年': '年', ' 日': '日', ' 月': '月',
    '、 ': '、', '： ': '：', ' ：': '：',
    ' 円': '円', ' 時': '時', ' 分': '分',
    ' 百万円': '百万円', ' 億円': '億円',
    ' ％': '％', ' %': '%',
    ',"': ',"', '  ': ' ',
    '",': '」,', ' 」': '」',
    '－': '–', '（ ': '（', ' （': '（',
    '） ': '）', '( ': '(', ' \"': ' "',
    '」 ': '」', ' 「': '「',
    '］ ': '］',
    'Source：': 'Source: ',
    '（Billions）': ' (Billions) ',
    'Note：': 'Note: ',
    '（Thousands）': '(Thousands) ',
    '（Millions）': ' (Millions) ',
    '〜 ': '〜',
    '\\t -': '\\t\u2013',
    '¥-': '¥–',
    '$ ': '$',
    '➢': '>',
    '\u3000\\t': '\\t',
    '\\\\s\\\\': '\\\\s¥',
    '\\\\s\\\\s\\\\s': '\\\\s',
    '\\t\\\\': '\\t¥',
}

# ── JA replacement dict ──
ja_replacement = {
    '  ': '', '\u00a0': '',
    '\u0021': '\uFF01', '\u0022': '\uFF02', '\u0023': '\uFF03',
    '\u0024': '\uFF04', '\uFF05': '\u0025', '\uFF06': '\u0026',
    '\u0027': '\uFF07', '\u0028': '\uFF08', '\u0029': '\uFF09',
    '\uFF0A': '\u002A', '\u002B': '\uFF0B',
    '\uFF0C': '\u002C',
    '\uFF0E': '\u002E',
    '\uFF10': '\u0030', '\uFF11': '\u0031', '\uFF12': '\u0032',
    '\uFF13': '\u0033', '\uFF14': '\u0034', '\uFF15': '\u0035',
    '\uFF16': '\u0036', '\uFF17': '\u0037', '\uFF18': '\u0038',
    '\uFF19': '\u0039',
    '\u003B': '\uFF1B', '\u003C': '\uFF1C', '\u003D': '\uFF1D',
    '\u003E': '\uFF1E', '\u003F': '\uFF1F', '\u0040': '\uFF20',
    '\uFF21': '\u0041', '\uFF22': '\u0042', '\uFF23': '\u0043',
    '\uFF24': '\u0044', '\uFF25': '\u0045', '\uFF26': '\u0046',
    '\uFF27': '\u0047', '\uFF28': '\u0048', '\uFF29': '\u0049',
    '\uFF2A': '\u004A', '\uFF2B': '\u004B', '\uFF2C': '\u004C',
    '\uFF2D': '\u004D', '\uFF2E': '\u004E', '\uFF2F': '\u004F',
    '\uFF30': '\u0050', '\uFF31': '\u0051', '\uFF32': '\u0052',
    '\uFF33': '\u0053', '\uFF34': '\u0054', '\uFF35': '\u0055',
    '\uFF36': '\u0056', '\uFF37': '\u0057', '\uFF38': '\u0058',
    '\uFF39': '\u0059', '\uFF3A': '\u005A',
    '\u005B': '\uFF3B', '\u005C': '\uFF3C', '\u005D': '\uFF3D',
    '\u005E': '\uFF3E', '\u005F': '\uFF3F', '\u0060': '\uFF40',
    '\uFF41': '\u0061', '\uFF42': '\u0062', '\uFF43': '\u0063',
    '\uFF44': '\u0064', '\uFF45': '\u0065', '\uFF46': '\u0066',
    '\uFF47': '\u0067', '\uFF48': '\u0068', '\uFF49': '\u0069',
    '\uFF4A': '\u006A', '\uFF4B': '\u006B', '\uFF4C': '\u006C',
    '\uFF4D': '\u006D', '\uFF4E': '\u006E', '\uFF4F': '\u006F',
    '\uFF50': '\u0070', '\uFF51': '\u0071', '\uFF52': '\u0072',
    '\uFF53': '\u0073', '\uFF54': '\u0074', '\uFF55': '\u0075',
    '\uFF56': '\u0076', '\uFF57': '\u0077', '\uFF58': '\u0078',
    '\uFF59': '\u0079', '\uFF5A': '\u007A',
    '\u007B': '\uFF5B', '\u007C': '\uFF5C', '\u007D': '\uFF5D',
    '\u007E': '\uFF5E',
    '\u33A1': 'm2', '\u33A5': 'm3',
    '\\\\ ': '¥', '("': '(「', '\\\\': '¥',
    "'": "'", '\" ': '"\\\\s', ' \"': '\\\\s「',
    '")': '」)', '\".': '」.', '.\"': '."',
    '・\\t': '• ', ' )': ')', ' . ': '. ',
    ' , ': ', ', '\" ': '」 ',
    '( ': '(', "' ": "' ", "'": "'",
    "• ": "•   ", '） ': '）', ' （': '（',
    '.\"': '."', "' ": "' ",
    ' 、': '、', '。 ': '。',
    ' 年': '年', ' 日': '日', ' 月': '月',
    '、 ': '、', '： ': '：', ' ：': '：',
    ' 円': '円', ' 時': '時', ' 分': '分',
    ' 百万円': '百万円', ' 億円': '億円',
    ' ％': '％', ' %': '%',
    ',"': ',"', '  ': ' ',
    '",': '」,', ' 」': '」',
    '（ ': '（', ' （': '（',
    '） ': '）', '( ': '(', ' \"': ' 「',
    '」 ': '」', ' 「': '「',
    '］ ': '］',
    'Source：': 'Source: ',
    '（Billions）': ' (Billions) ',
    'Note：': 'Note: ',
    '（Thousands）': '(Thousands) ',
    '（Millions）': ' (Millions) ',
    '〜 ': '〜',
    '\\t -': '\\t\u2013',
    '¥-': '¥–',
    '$ ': '$',
    '➢': '>',
    '\u3000\\t': '\\t',
    '\\\\s\\\\': '\\\\s¥',
    '\\\\s\\\\s\\\\s': '\\\\s',
    '\\t\\\\': '\\t¥',
}

# JA-specific helpers
_url_re = re.compile(r'https?[:\\uFF1A][/\\uFF0F]{2}[^\s\u3000]+')

def protect_urls(text, func):
    urls = []
    def save(m):
        url = m.group(0).replace('\uFF0F', '/').replace('\uFF1A', ':')
        urls.append(url)
        return f'\x00URL{len(urls)-1}\x00'
    text = _url_re.sub(save, text)
    text = func(text)
    for i, url in enumerate(urls):
        text = text.replace(f'\x00URL{i}\x00', url)
    return text

def apply_slash_colon(text):
    text = text.replace('/', '\uff0f')
    text = text.replace(':', '\uff1a')
    return text

def strip_line_head_spaces(text):
    return re.sub(r'(^|\r?\n)[ \u00A0\u3000\t]+', r'\1', text)

def strip_spaces_around_numbers(text):
    protected_spaces = []
    def protect_note_space(match):
        protected_spaces.append(match.group(0))
        return f'\x00NOTE_SPACE{len(protected_spaces) - 1}\x00'
    text = re.sub(r'([\u203B\uFF0A\*][0-9\uFF10-\uFF19]+) ', protect_note_space, text)
    text = re.sub(r'\s*(\d+)\s*', r'\1', text)
    for index, protected in enumerate(protected_spaces):
        text = text.replace(f'\x00NOTE_SPACE{index}\x00', protected)
    return text



# ── process functions ──

def process_en_pptx(input_path, output_path):
    """Apply English normalization to a PPTX file."""
    prs = Presentation(input_path)
    for slide in prs.slides:
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            for para in shape.text_frame.paragraphs:
                for run in para.runs:
                    t = run.text
                    t = normalize_url_email(t)
                    for old, new in en_replacement.items():
                        t = t.replace(old, new)
                    t = re.sub('(?<!\\s)\\(', ' (', t)
                    t = re.sub('\\)(?!\\s|.|,)', ') ', t)
                    t = re.sub(':(?![/\\s])', ': ', t)
                    t = re.sub('[ ]+(\r?\n)', '\1', t)
                    run.text = t
                # 段落末尾の空白/NBSP削除（改行前の半角スペース対応）
                ns = '{http://schemas.openxmlformats.org/drawingml/2006/main}'
                _last_t = None
                for _elem in para._p.iter():
                    if _elem.tag == f'{ns}t' and _elem.text:
                        _last_t = _elem
                    elif _elem.tag == f'{ns}endParaRPr' and _last_t is not None:
                        stripped = _last_t.text.rstrip(' \t\xa0')
                        if stripped != _last_t.text:
                            _last_t.text = stripped
                if _last_t is not None:
                    stripped = _last_t.text.rstrip(' \t\xa0')
                    if stripped != _last_t.text:
                        _last_t.text = stripped
    prs.save(output_path)
    print(f"  EN processed: {input_path} → {output_path}")

def process_ja_pptx(input_path, output_path):
    """Apply Japanese normalization to a PPTX file."""
    prs = Presentation(input_path)
    for slide in prs.slides:
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            for para in shape.text_frame.paragraphs:
                for run in para.runs:
                    t = run.text
                    for old, new in ja_replacement.items():
                        t = t.replace(old, new)
                    run.text = t
                if para.runs:
                    full_text = ''.join(run.text for run in para.runs)
                    full_text = normalize_url_email(full_text)
                    full_text = protect_urls(full_text, apply_slash_colon)
                    full_text = strip_spaces_around_numbers(full_text)
                    full_text = strip_line_head_spaces(full_text)
                    idx = 0
                    for run in para.runs:
                        run_len = len(run.text)
                        run.text = full_text[idx:idx + run_len]
                        idx += run_len
                pass
                pass
    prs.save(output_path)
    print(f"  JA processed: {input_path} → {output_path}")

def extract_text_summary(path, label):
    """Read back the PPTX text for verification."""
    prs = Presentation(path)
    lines = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            for para in shape.text_frame.paragraphs:
                text = para.text
                if text.strip():
                    lines.append(text)
    print(f"\n── {label} ──")
    for line in lines[:30]:
        print(f"  | {line}")
    if len(lines) > 30:
        print(f"  | ... ({len(lines) - 30} more lines)")
    return lines


# ── main ──
if __name__ == "__main__":
    tests_dir = os.path.dirname(__file__)
    output_dir = os.path.join(tests_dir, "..", "output")
    os.makedirs(output_dir, exist_ok=True)

    en_input = os.path.join(tests_dir, "test_en.pptx")
    en_output = os.path.join(output_dir, "test_en_chk.pptx")
    ja_input = os.path.join(tests_dir, "test_ja.pptx")
    ja_output = os.path.join(output_dir, "test_ja_chk.pptx")

    print("=" * 60)
    print("Processing English normalization test...")
    process_en_pptx(en_input, en_output)

    print("\nProcessing Japanese normalization test...")
    process_ja_pptx(ja_input, ja_output)

    # Read back results
    extract_text_summary(en_output, "EN output")
    extract_text_summary(ja_output, "JA output")

    print("\n✅ Done. Output files:")
    print(f"  {en_output}")
    print(f"  {ja_output}")
